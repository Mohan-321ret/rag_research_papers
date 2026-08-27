"""Text and metadata extraction for each supported file type."""

import io
from dataclasses import dataclass, field

from app.core.exceptions import BadRequestError
from app.core.logging import get_logger
from app.modules.ingestion.detection import FileType
from app.modules.ingestion.ocr import TesseractOcrEngine

logger = get_logger(__name__)

# Below this average per-page character count a PDF is treated as scanned
# and sent to OCR (when available).
_OCR_CHARS_PER_PAGE_THRESHOLD = 20


@dataclass(slots=True)
class ExtractionResult:
    """Extracted text plus whatever native metadata the format carries."""

    text: str
    native_metadata: dict[str, str] = field(default_factory=dict)
    page_count: int | None = None
    used_ocr: bool = False
    warnings: list[str] = field(default_factory=list)


def _clean(value: object) -> str | None:
    text = str(value).strip() if value is not None else ""
    return text or None


class PdfExtractor:
    def __init__(self, ocr_engine: TesseractOcrEngine | None) -> None:
        self._ocr = ocr_engine

    def extract(self, content: bytes) -> ExtractionResult:
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception as exc:
            raise BadRequestError(f"Could not parse PDF file: {exc}") from exc

        # Pages joined with \f so the processing pipeline can assign
        # page numbers to chunks (for citations).
        text = "\f".join((page.extract_text() or "") for page in reader.pages).strip("\f").strip()
        result = ExtractionResult(text=text, page_count=len(reader.pages))

        meta = reader.metadata
        if meta is not None:
            result.native_metadata = {
                key: value
                for key, value in {
                    "title": _clean(meta.title),
                    "author": _clean(meta.author),
                    "subject": _clean(meta.subject),
                }.items()
                if value is not None
            }

        pages = max(len(reader.pages), 1)
        if len(text) / pages < _OCR_CHARS_PER_PAGE_THRESHOLD:
            result = self._try_ocr(content, result)
        return result

    def _try_ocr(self, content: bytes, result: ExtractionResult) -> ExtractionResult:
        if self._ocr is None:
            result.warnings.append(
                "PDF appears to be scanned but OCR is disabled or unavailable"
            )
            return result
        logger.info("pdf_ocr_fallback", pages=result.page_count)
        try:
            ocr_text = self._ocr.extract_pdf_text(content)
        except Exception as exc:
            result.warnings.append(f"OCR failed: {type(exc).__name__}")
            return result
        if len(ocr_text) > len(result.text):
            result.text = ocr_text
            result.used_ocr = True
        return result


class DocxExtractor:
    def extract(self, content: bytes) -> ExtractionResult:
        from docx import Document as DocxDocument

        try:
            document = DocxDocument(io.BytesIO(content))
        except Exception as exc:
            raise BadRequestError(f"Could not parse DOCX file: {exc}") from exc

        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append("\t".join(cells))

        props = document.core_properties
        native = {
            key: value
            for key, value in {
                "title": _clean(props.title),
                "author": _clean(props.author),
                "subject": _clean(props.subject),
            }.items()
            if value is not None
        }
        return ExtractionResult(text="\n".join(parts).strip(), native_metadata=native)


class TxtExtractor:
    def extract(self, content: bytes) -> ExtractionResult:
        from charset_normalizer import from_bytes

        best = from_bytes(content).best()
        if best is None:
            raise BadRequestError("Could not decode text file (unknown encoding)")
        result = ExtractionResult(text=str(best).strip())
        if best.encoding not in ("utf_8", "ascii"):
            result.warnings.append(f"Decoded with detected encoding {best.encoding}")
        return result


class HtmlExtractor:
    def extract(self, content: bytes) -> ExtractionResult:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(content, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()

        native: dict[str, str] = {}
        if soup.title is not None and _clean(soup.title.string):
            native["title"] = _clean(soup.title.string)  # type: ignore[assignment]
        author_meta = soup.find("meta", attrs={"name": "author"})
        if author_meta is not None and _clean(author_meta.get("content")):
            native["author"] = _clean(author_meta.get("content"))  # type: ignore[assignment]

        text = "\n".join(
            line.strip() for line in soup.get_text("\n").splitlines() if line.strip()
        )
        return ExtractionResult(text=text, native_metadata=native)


class XlsxExtractor:
    def extract(self, content: bytes) -> ExtractionResult:
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:
            raise BadRequestError(f"Could not parse XLSX file: {exc}") from exc

        try:
            sheets: list[str] = []
            for sheet in workbook.worksheets:
                rows = [
                    "\t".join("" if cell is None else str(cell) for cell in row)
                    for row in sheet.iter_rows(values_only=True)
                    if any(cell is not None for cell in row)
                ]
                if rows:
                    sheets.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))

            props = workbook.properties
            native = {
                key: value
                for key, value in {
                    "title": _clean(props.title),
                    "author": _clean(props.creator),
                }.items()
                if value is not None
            }
            return ExtractionResult(
                text="\n\n".join(sheets).strip(),
                native_metadata=native,
                page_count=len(workbook.worksheets),
            )
        finally:
            workbook.close()


def build_extractor_registry(ocr_engine: TesseractOcrEngine | None):
    """Map each supported file type to its extractor."""
    return {
        FileType.PDF: PdfExtractor(ocr_engine),
        FileType.DOCX: DocxExtractor(),
        FileType.TXT: TxtExtractor(),
        FileType.HTML: HtmlExtractor(),
        FileType.XLSX: XlsxExtractor(),
    }
