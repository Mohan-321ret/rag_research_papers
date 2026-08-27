"""Tests for the /api/v1/documents endpoints."""

import io
import uuid

from fastapi.testclient import TestClient

TXT_CONTENT = b"Drift-aware adaptive retrieval augmented generation for enterprises."

HTML_CONTENT = (
    b"<!doctype html><html><head><title>Drift Survey</title>"
    b'<meta name="author" content="A. Researcher"></head>'
    b"<body><h1>Knowledge Drift</h1><p>Enterprise knowledge changes.</p>"
    b"<script>ignore_me()</script></body></html>"
)


def _docx_bytes() -> bytes:
    from docx import Document

    buffer = io.BytesIO()
    document = Document()
    document.add_paragraph("DAA-RAG design notes.")
    document.add_paragraph("Versioning enables temporal queries.")
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "metrics"
    sheet.append(["metric", "value"])
    sheet.append(["drift_score", 0.42])
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, headers, filename: str, content: bytes, **form):
    return client.post(
        "/api/v1/documents/upload",
        headers=headers,
        files={"file": (filename, content, "application/octet-stream")},
        data=form,
    )


def test_upload_txt(client: TestClient, auth_headers) -> None:
    response = _upload(client, auth_headers, "notes.txt", TXT_CONTENT, department="R&D")

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["status"] == "READY"
    assert body["current_version"] == 1
    assert body["new_version_created"] is True
    metadata = body["metadata"]
    assert metadata["file_type"] == "txt"
    assert metadata["filename"] == "notes.txt"
    assert metadata["title"] == "notes"
    assert metadata["source"] == "upload"
    assert metadata["department"] == "R&D"
    assert metadata["access_level"] == "internal"
    assert metadata["characters"] == len(TXT_CONTENT.decode())


def test_reupload_identical_content_is_deduplicated(client: TestClient, auth_headers) -> None:
    first = _upload(client, auth_headers, "dedup.txt", TXT_CONTENT).json()
    second = _upload(client, auth_headers, "dedup.txt", TXT_CONTENT).json()

    assert second["id"] == first["id"]
    assert second["deduplicated"] is True
    assert second["new_version_created"] is False
    assert second["version_count"] == 1


def test_reupload_changed_content_creates_new_version(
    client: TestClient, auth_headers
) -> None:
    first = _upload(client, auth_headers, "evolving.txt", b"Policy v1: remote work allowed.")
    document_id = first.json()["id"]

    second = _upload(client, auth_headers, "evolving.txt", b"Policy v2: hybrid only.")
    assert second.json()["current_version"] == 2
    assert second.json()["version_count"] == 2

    versions = client.get(
        f"/api/v1/documents/{document_id}/versions", headers=auth_headers
    ).json()
    assert [v["version_number"] for v in versions] == [1, 2]
    assert [v["is_current"] for v in versions] == [False, True]
    assert versions[0]["content_hash"] != versions[1]["content_hash"]


def test_upload_html_uses_native_title_and_author(client: TestClient, auth_headers) -> None:
    body = _upload(client, auth_headers, "survey.html", HTML_CONTENT).json()
    assert body["metadata"]["file_type"] == "html"
    assert body["metadata"]["title"] == "Drift Survey"
    assert body["metadata"]["author"] == "A. Researcher"
    assert body["title"] == "Drift Survey"


def test_upload_docx(client: TestClient, auth_headers) -> None:
    body = _upload(
        client, auth_headers, "design.docx", _docx_bytes(), author="B. Writer"
    ).json()
    assert body["metadata"]["file_type"] == "docx"
    assert body["metadata"]["author"] == "B. Writer"


def test_upload_xlsx(client: TestClient, auth_headers) -> None:
    body = _upload(client, auth_headers, "metrics.xlsx", _xlsx_bytes()).json()
    assert body["metadata"]["file_type"] == "xlsx"
    assert body["metadata"]["page_count"] == 1


def test_upload_pdf(client: TestClient, auth_headers, minimal_pdf_factory) -> None:
    pdf = minimal_pdf_factory("Adaptive retrieval handles knowledge drift gracefully.")
    body = _upload(client, auth_headers, "paper.pdf", pdf).json()
    assert body["metadata"]["file_type"] == "pdf"
    assert body["metadata"]["page_count"] == 1
    assert body["status"] == "READY"


def test_upload_unsupported_type_rejected(client: TestClient, auth_headers) -> None:
    response = _upload(client, auth_headers, "malware.exe", b"\x00\x01\x02\x03")
    assert response.status_code == 415
    assert "Unsupported" in response.json()["detail"]


def test_upload_empty_file_rejected(client: TestClient, auth_headers) -> None:
    assert _upload(client, auth_headers, "empty.txt", b"").status_code == 400


def test_upload_requires_auth(client: TestClient) -> None:
    response = client.post(
        "/api/v1/documents/upload", files={"file": ("a.txt", b"hello", "text/plain")}
    )
    assert response.status_code == 401


def test_list_and_search(client: TestClient, auth_headers) -> None:
    marker = uuid.uuid4().hex[:8]
    _upload(client, auth_headers, f"report-{marker}.txt", TXT_CONTENT, title=f"Report {marker}")

    listing = client.get("/api/v1/documents", headers=auth_headers).json()
    assert listing["total"] >= 1

    found = client.get(
        f"/api/v1/documents?search={marker}", headers=auth_headers
    ).json()
    assert found["total"] == 1
    assert found["items"][0]["title"] == f"Report {marker}"


def test_get_unknown_document_404(client: TestClient, auth_headers) -> None:
    response = client.get(f"/api/v1/documents/{uuid.uuid4()}", headers=auth_headers)
    assert response.status_code == 404


def test_delete_by_owner(client: TestClient, auth_headers) -> None:
    document_id = _upload(client, auth_headers, "todelete.txt", TXT_CONTENT).json()["id"]

    assert (
        client.delete(f"/api/v1/documents/{document_id}", headers=auth_headers).status_code
        == 204
    )
    assert (
        client.get(f"/api/v1/documents/{document_id}", headers=auth_headers).status_code
        == 404
    )


def test_delete_by_other_user_forbidden(client: TestClient, auth_headers) -> None:
    document_id = _upload(client, auth_headers, "mine.txt", TXT_CONTENT).json()["id"]

    other = client.post(
        "/api/v1/auth/register",
        json={"email": f"other-{uuid.uuid4().hex[:10]}@example.com", "password": "secret-password-1"},
    ).json()
    other_headers = {"Authorization": f"Bearer {other['access_token']}"}

    response = client.delete(f"/api/v1/documents/{document_id}", headers=other_headers)
    assert response.status_code == 403
